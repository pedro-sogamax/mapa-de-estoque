"""Testes da revisao semanal: o relatorio a revisora e a conferencia da leitura.

O que mais importa, nesta ordem:

* a conferencia NUNCA marca a mensagem como lida — senao ela forjaria a propria prova;
* "nao consegui conferir" nunca vira "sem leitura", e o alerta de prazo sai uma vez so;
* o relatorio da semana sai uma vez, e se nao saiu na segunda, sai na rodada seguinte.
"""

from __future__ import annotations

import json
from datetime import date, datetime
from pathlib import Path

import pytest
from openpyxl import load_workbook

from src import alerta, historico, revisao
from src.config import Config
from src.disparo.canais.base import Envio, FalhaNoEnvio
from src.disparo.eml import montar_mensagem
from src.revisao import (
    AGUARDANDO,
    LIDO,
    LIDO_APOS_O_PRAZO,
    SEM_LEITURA,
    FalhaNaConferencia,
    Registro,
    Relatorio,
    conferir,
    dias_uteis_desde,
    enviar_relatorio,
    montar_relatorio,
    semana_de,
)
from src.trava import Trava

SEGUNDA = date(2026, 9, 28)


def _cfg(tmp_path: Path, **troca) -> Config:
    base = dict(
        url_base="https://geweb", usuario="u", senha="s",
        download_dir=tmp_path / "downloads", formatado_dir=tmp_path / "formatado",
        envios_dir=tmp_path / "envios", templates_dir=tmp_path / "templates",
        auth_state_path=tmp_path / "state.json", sequencia_path=tmp_path / "sequencia.json",
        estado_path=tmp_path / "estado.json", rodada_path=tmp_path / "rodada.json",
        envios_path=tmp_path / "envios.json", envios_historico_path=tmp_path / "envios.jsonl",
        historico_path=tmp_path / "historico.jsonl", historico_dir=tmp_path / "formatado" / "logs",
        headless=True, timeout_ms=1, timeout_relatorio_ms=1, slow_mo_ms=0,
        comprador="", comprador_cargo="", comprador_telefones="",
        remetente="", responder_para="", destinatario_teste=(), alerta_para=("ti@x.com",),
        telefone_teste="", smtp_host="smtp.x", smtp_porta=465, smtp_usuario="confirmacao@x.com",
        smtp_senha="s", smtp_seguranca="ssl", whatsapp_provedor="zapi", whatsapp_url_base="",
        whatsapp_instancia="", whatsapp_token="", whatsapp_client_token="",
        intervalo_envio_s=0, max_envios_por_rodada=95, max_envios_por_hora=90,
        max_falhas_seguidas=3, max_tentativas=1, max_anexo_mb=10,
        revisoes_path=tmp_path / "dados" / "revisoes.json",
        revisoes_dir=tmp_path / "logs" / "revisoes",
        trava_path=tmp_path / "dados" / "rodada.trava",
        revisao_para=("revisora@x.com",),
        revisao_copia_para=("ti@x.com",),
        revisao_alerta_para=("gerente@x.com",),
        revisao_prazo_dias=3,
        imap_host="imap.x",
        revisao_imap_usuario="revisora@x.com",
        revisao_imap_senha="s",
    )
    base.update(troca)
    return Config(**base)


def _eventos_da_semana():
    def em(dia: str) -> dict:
        return {"em": f"{dia}T07:10:00"}

    return [
        {**historico.evento_extracao("Yuri", "MARJAN", "2026-09", "ok"), **em("2026-09-23")},
        {**historico.evento_extracao("Yuri", "ACHE", "2026-09", "falha", "timeout no Geweb"), **em("2026-09-24")},
        {**historico.evento_envio("Yuri", "MARJAN", "2026-09", "email", "a@b", "ok"), **em("2026-09-24")},
        {**historico.evento_envio("Yuri", "LEBON", "2026-09", "email", "c@d", "nao tentado", "cota atingida"), **em("2026-09-28")},
        # fora da semana (terca 22/09 a segunda 28/09): nao pode aparecer
        {**historico.evento_extracao("Yuri", "ANTIGO", "2026-09", "falha", "velho"), **em("2026-09-21")},
    ]


class CanalFalso:
    enviados: list = []

    def __init__(self, cfg):
        pass

    def __enter__(self):
        return self

    def __exit__(self, *_):
        return False

    def enviar(self, destino, mensagem, anexo):
        CanalFalso.enviados.append((destino, mensagem, anexo))
        return Envio(canal="email", destinatarios=destino.emails, identificadores=(f"<id{len(CanalFalso.enviados)}@x>",))


@pytest.fixture
def canal(monkeypatch):
    CanalFalso.enviados = []
    monkeypatch.setattr(revisao, "CanalEmail", CanalFalso)
    return CanalFalso


@pytest.fixture
def alertas(monkeypatch):
    enviados = []

    def falso(cfg, assunto, linhas, destinatarios=None, rodape_tecnico=True):
        enviados.append({"assunto": assunto, "texto": "\n".join(linhas), "para": destinatarios})
        return True

    monkeypatch.setattr(alerta, "enviar", falso)
    return enviados


def _registrar(cfg: Config, **troca) -> Relatorio:
    relatorio = Relatorio(
        inicio="2026-09-22", fim="2026-09-28", message_id="<id1@x>",
        enviado_em="2026-09-28T07:15:00", para=["revisora@x.com"],
    )
    for campo, valor in troca.items():
        setattr(relatorio, campo, valor)
    registro = Registro(cfg.revisoes_path)
    registro.relatorios.append(relatorio)
    registro.gravar()
    return relatorio


def _situacao(cfg: Config) -> Relatorio:
    return Registro(cfg.revisoes_path).ultimo()


class TestDatas:
    def test_a_semana_vai_de_terca_a_segunda(self):
        assert semana_de(SEGUNDA) == (date(2026, 9, 22), SEGUNDA)

    def test_depois_da_segunda_continua_sendo_a_semana_que_fechou_nela(self):
        assert semana_de(date(2026, 10, 2)) == (date(2026, 9, 22), SEGUNDA)

    def test_segunda_a_quinta_sao_tres_dias_uteis(self):
        assert dias_uteis_desde(SEGUNDA, date(2026, 10, 1)) == 3

    def test_feriado_nao_conta(self):
        """12/10/2026 e segunda e feriado: relatorio da sexta 09/10 vence na quinta 15/10."""
        assert dias_uteis_desde(date(2026, 10, 9), date(2026, 10, 14)) == 2
        assert dias_uteis_desde(date(2026, 10, 9), date(2026, 10, 15)) == 3


class TestTextoDoRelatorio:
    def test_conta_e_lista_as_pendencias_da_semana(self):
        eventos = revisao.eventos_da_semana(_eventos_da_semana(), date(2026, 9, 22), SEGUNDA)
        assunto, corpo = montar_relatorio(eventos, date(2026, 9, 22), SEGUNDA)

        assert "22/09 a 28/09/2026" in assunto and "2 pendência(s)" in assunto
        assert "Extrações: 1 ok, 1 falha(s)" in corpo
        assert "ACHE" in corpo and "timeout no Geweb" in corpo
        assert "LEBON" in corpo and "[não tentado]" in corpo
        assert "ANTIGO" not in corpo

    def test_semana_limpa(self):
        eventos = [{**historico.evento_extracao("Y", "MARJAN", "2026-09", "ok"), "em": "2026-09-23T07:00:00"}]
        assunto, corpo = montar_relatorio(eventos, date(2026, 9, 22), SEGUNDA)
        assert "pendência" not in assunto
        assert "Nenhuma falha na semana." in corpo

    def test_rodada_quebrada_entra_no_topo(self):
        _, corpo = montar_relatorio([], date(2026, 9, 22), SEGUNDA, erro_da_rodada="login recusado")
        assert "foi interrompida" in corpo and "login recusado" in corpo

    def test_avisa_que_a_leitura_e_registrada(self):
        """Transparencia: quem recebe sabe que a abertura fica registrada."""
        _, corpo = montar_relatorio([], date(2026, 9, 22), SEGUNDA)
        assert "abertura deste e-mail fica registrada" in corpo


class TestEnvioDoRelatorio:
    def test_envia_com_a_planilha_da_semana_e_registra(self, tmp_path, canal):
        cfg = _cfg(tmp_path)
        historico.registrar(cfg.historico_path, _eventos_da_semana())

        assert enviar_relatorio(cfg, hoje=SEGUNDA) is True

        destino, mensagem, anexo = canal.enviados[0]
        assert destino.emails == ("revisora@x.com",) and destino.copia_oculta == ("ti@x.com",)
        assert anexo is not None and anexo.exists()
        extracoes = load_workbook(anexo)["Extracoes"]
        assert extracoes.max_row == 3  # cabecalho + MARJAN + ACHE; o de 21/09 ficou de fora
        relatorio = _situacao(cfg)
        assert (relatorio.fim, relatorio.message_id, relatorio.situacao) == ("2026-09-28", "<id1@x>", AGUARDANDO)

    def test_nao_reenvia_a_mesma_semana(self, tmp_path, canal):
        cfg = _cfg(tmp_path)
        enviar_relatorio(cfg, hoje=SEGUNDA)

        assert enviar_relatorio(cfg, hoje=SEGUNDA) is False
        assert enviar_relatorio(cfg, hoje=date(2026, 9, 30)) is False
        assert len(canal.enviados) == 1

    def test_semana_que_nao_saiu_na_segunda_sai_na_rodada_seguinte(self, tmp_path, canal):
        cfg = _cfg(tmp_path)

        assert enviar_relatorio(cfg, hoje=date(2026, 9, 30)) is True
        assert _situacao(cfg).fim == "2026-09-28"

    def test_erro_da_rodada_so_entra_na_propria_segunda(self, tmp_path, canal):
        cfg = _cfg(tmp_path)
        enviar_relatorio(cfg, hoje=date(2026, 9, 29), erro_da_rodada="login recusado")
        assert "login recusado" not in canal.enviados[0][1].corpo

    def test_falha_no_smtp_nao_registra_e_fica_para_a_proxima(self, tmp_path, monkeypatch):
        class Quebrado(CanalFalso):
            def __enter__(self):
                raise FalhaNoEnvio("servidor fora")

        monkeypatch.setattr(revisao, "CanalEmail", Quebrado)
        cfg = _cfg(tmp_path)

        assert enviar_relatorio(cfg, hoje=SEGUNDA) is False
        assert Registro(cfg.revisoes_path).relatorios == []

    def test_desligada_sem_revisora(self, tmp_path, canal):
        assert enviar_relatorio(_cfg(tmp_path, revisao_para=()), hoje=SEGUNDA) is False
        assert canal.enviados == []


class TestConferencia:
    def test_sem_relatorio_nao_conecta(self, tmp_path):
        def nao_pode(*_):
            raise AssertionError("conectou sem relatorio pendente")

        assert conferir(_cfg(tmp_path), verificar=nao_pode) is False

    def test_lido_registra_a_hora_da_conferencia(self, tmp_path):
        cfg = _cfg(tmp_path)
        _registrar(cfg)

        assert conferir(cfg, datetime(2026, 9, 28, 11, 0), verificar=lambda *_: True) is True

        relatorio = _situacao(cfg)
        assert relatorio.situacao == LIDO and relatorio.lido_ate == "2026-09-28T11:00:00"

    def test_lido_nao_e_conferido_de_novo(self, tmp_path):
        cfg = _cfg(tmp_path)
        _registrar(cfg, situacao=LIDO, lido_ate="2026-09-28T11:00:00")

        def nao_pode(*_):
            raise AssertionError("conferiu um relatorio ja lido")

        assert conferir(cfg, datetime(2026, 9, 29, 9, 0), verificar=nao_pode) is False

    def test_nao_lido_dentro_do_prazo_nao_muda_nada(self, tmp_path, alertas):
        cfg = _cfg(tmp_path)
        _registrar(cfg)

        assert conferir(cfg, datetime(2026, 9, 30, 17, 0), verificar=lambda *_: False) is False
        assert _situacao(cfg).situacao == AGUARDANDO and alertas == []

    def test_prazo_vencido_marca_sem_leitura_e_avisa_a_gerente_uma_vez(self, tmp_path, alertas):
        cfg = _cfg(tmp_path)
        _registrar(cfg)

        conferir(cfg, datetime(2026, 10, 1, 8, 0), verificar=lambda *_: False)
        conferir(cfg, datetime(2026, 10, 1, 9, 0), verificar=lambda *_: False)

        assert _situacao(cfg).situacao == SEM_LEITURA
        assert len(alertas) == 1 and alertas[0]["para"] == ("gerente@x.com",)
        assert "22/09 a 28/09/2026" in alertas[0]["assunto"]

    def test_lido_depois_do_prazo(self, tmp_path):
        cfg = _cfg(tmp_path)
        _registrar(cfg, situacao=SEM_LEITURA, alerta_em="2026-10-01T08:00:00")

        conferir(cfg, datetime(2026, 10, 2, 10, 0), verificar=lambda *_: True)

        assert _situacao(cfg).situacao == LIDO_APOS_O_PRAZO

    def test_alerta_que_nao_saiu_e_tentado_de_novo(self, tmp_path, monkeypatch):
        tentativas = []
        monkeypatch.setattr(alerta, "enviar", lambda *a, **k: tentativas.append(1) or len(tentativas) > 1)
        cfg = _cfg(tmp_path)
        _registrar(cfg)

        conferir(cfg, datetime(2026, 10, 1, 8, 0), verificar=lambda *_: False)
        assert _situacao(cfg).alerta_em == ""
        conferir(cfg, datetime(2026, 10, 1, 9, 0), verificar=lambda *_: False)
        assert _situacao(cfg).alerta_em == "2026-10-01T09:00:00" and len(tentativas) == 2

    def test_mensagem_sumida_avisa_que_pode_ter_sido_apagada(self, tmp_path, alertas):
        cfg = _cfg(tmp_path)
        _registrar(cfg)

        conferir(cfg, datetime(2026, 10, 1, 8, 0), verificar=lambda *_: None)

        assert "pode ter sido apagada" in alertas[0]["texto"]

    def test_falha_tecnica_nao_vira_sem_leitura_e_avisa_a_ti_uma_vez_por_dia(self, tmp_path, alertas):
        cfg = _cfg(tmp_path)
        _registrar(cfg)

        def quebrado(*_):
            raise FalhaNaConferencia("senha recusada")

        conferir(cfg, datetime(2026, 10, 1, 8, 0), verificar=quebrado)
        conferir(cfg, datetime(2026, 10, 1, 9, 0), verificar=quebrado)
        conferir(cfg, datetime(2026, 10, 2, 8, 0), verificar=quebrado)

        assert _situacao(cfg).situacao == AGUARDANDO
        assert len(alertas) == 2 and alertas[0]["para"] is None  # None = ALERTA_PARA

    def test_a_aba_revisoes_entra_na_planilha_do_mes(self, tmp_path):
        cfg = _cfg(tmp_path)
        historico.registrar(cfg.historico_path, _eventos_da_semana())
        _registrar(cfg)

        conferir(cfg, datetime(2026, 9, 28, 11, 0), verificar=lambda *_: True)

        wb = load_workbook(cfg.historico_dir / "2026-09.xlsx")
        linhas = list(wb["Revisoes"].iter_rows(values_only=True))
        assert linhas[0] == ("SEMANA", "ENVIADO EM", "LIDO ATÉ", "SITUAÇÃO")
        assert linhas[1] == ("22/09 a 28/09/2026", "28/09/2026 07:15", "28/09/2026 11:00", LIDO)


class ImapFalso:
    """Caixa com duas pastas. Registra cada comando para o teste conferir que so houve leitura."""

    comandos: list = []

    def __init__(self, host, porta, timeout=None):
        self.comandos.append(("conectar", host))

    def login(self, usuario, senha):
        self.comandos.append(("login", usuario))

    def list(self):
        return "OK", [b'(\\HasChildren) "." INBOX', b'(\\HasNoChildren \\Junk) "." INBOX.Mala_Direta']

    def select(self, pasta, readonly=False):
        self.comandos.append(("select", pasta, readonly))
        self.pasta = pasta
        return "OK", [b"1"]

    def search(self, charset, *criterio):
        self.comandos.append(("search", criterio))
        return "OK", [b"7" if self.pasta == '"INBOX.Mala_Direta"' else b""]

    def fetch(self, numero, partes):
        self.comandos.append(("fetch", partes))
        return "OK", [b"7 (FLAGS (\\Seen))"]

    def logout(self):
        self.comandos.append(("logout",))


class TestLeituraPorImap:
    @pytest.fixture(autouse=True)
    def imap(self, monkeypatch):
        ImapFalso.comandos = []
        monkeypatch.setattr(revisao.imaplib, "IMAP4_SSL", ImapFalso)

    def test_acha_a_mensagem_em_qualquer_pasta(self, tmp_path):
        assert revisao.verificar_leitura(_cfg(tmp_path), "<id1@x>") is True

    def test_so_abre_pasta_em_modo_leitura_e_so_le_marcadores(self, tmp_path):
        """O invariante central: SELECT ou FETCH do corpo marcariam o relatorio como lido."""
        revisao.verificar_leitura(_cfg(tmp_path), "<id1@x>")

        selects = [c for c in ImapFalso.comandos if c[0] == "select"]
        assert selects and all(readonly for _, _, readonly in selects)
        assert all(c[1] == "(FLAGS)" for c in ImapFalso.comandos if c[0] == "fetch")
        assert ImapFalso.comandos[-1] == ("logout",)

    def test_sem_credencial_e_falha_de_conferencia(self, tmp_path):
        with pytest.raises(FalhaNaConferencia):
            revisao.verificar_leitura(_cfg(tmp_path, revisao_imap_senha=""), "<id1@x>")


class TestOrdemAoFimDaRodada:
    def test_confere_o_anterior_antes_de_enviar_o_novo(self, tmp_path, monkeypatch):
        ordem = []
        monkeypatch.setattr(revisao, "conferir", lambda cfg: ordem.append("conferir"))
        monkeypatch.setattr(revisao, "enviar_relatorio", lambda cfg, erro_da_rodada="": ordem.append("enviar"))

        revisao.ao_fim_da_rodada(_cfg(tmp_path))

        assert ordem == ["conferir", "enviar"]

    def test_nunca_derruba_a_rodada(self, tmp_path, monkeypatch):
        def explode(cfg):
            raise RuntimeError("inesperado")

        monkeypatch.setattr(revisao, "conferir", explode)
        revisao.ao_fim_da_rodada(_cfg(tmp_path))  # nao pode levantar


class TestRegistro:
    def test_arquivo_ilegivel_e_guardado_em_vez_de_sobrescrito(self, tmp_path):
        arquivo = tmp_path / "revisoes.json"
        arquivo.write_text("{quebrado", encoding="utf-8")

        registro = Registro(arquivo)
        registro.gravar()

        assert (tmp_path / "revisoes.json.ilegivel").read_text(encoding="utf-8") == "{quebrado"
        assert json.loads(arquivo.read_text(encoding="utf-8"))["relatorios"] == []


class TestCopiaOculta:
    def test_vai_no_bcc_e_nao_no_cc(self):
        msg = montar_mensagem(["a@x"], [], "assunto", "corpo", None, rascunho=False, copia_oculta=["b@x"])
        assert msg["Bcc"] == "b@x" and msg["Cc"] is None


class TestTrava:
    def test_segunda_execucao_nao_entra_enquanto_a_primeira_segura(self, tmp_path):
        primeira, segunda = Trava(tmp_path / "t.trava"), Trava(tmp_path / "t.trava")

        assert primeira.adquirir() is True
        assert segunda.adquirir() is False
        primeira.liberar()
        assert segunda.adquirir() is True
        segunda.liberar()
