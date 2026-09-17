"""Testes do historico: o jsonl e a planilha que os compradores abrem.

Dois lados precisam continuar valendo, e os dois ja custaram uma decisao de desenho:

* a planilha vive DENTRO do FORMATADO_DIR, que o disparo varre com rglob e a sequencia
  varre atras de numero. Se o nome for confundido, o `--periodo` monta uma leva com um
  arquivo de log no lugar do mapa, ou o contador pula numeros;
* gravar o log nunca pode derrubar uma rodada que ja terminou -- o trabalho ja foi feito.
"""

from __future__ import annotations

from pathlib import Path

from openpyxl import load_workbook

from src import historico
from src.disparo.selecao import _rotulo_do_arquivo
from src.sequencia import _PADRAO_NUMERO


def _eventos():
    return [
        historico.evento_extracao("Yuri Toso", "MARJAN", "2026-09", "ok", arquivo="0112_x.xlsx"),
        historico.evento_extracao("Joici Rangel", "ABBOTT", "2026-09", "falha", "timeout"),
        historico.evento_envio("Joici Rangel", "ABBOTT", "2026-09", "email", "a@b.com", "ok",
                               identificador="<id@x>"),
        historico.evento_envio("Joici Rangel", "LEBON", "2026-09", "email", "c@d.com",
                               "nao tentado", "cota atingida"),
    ]


class TestRegistroEmJsonl:
    def test_grava_e_le_de_volta(self, tmp_path):
        arquivo = tmp_path / "historico.jsonl"
        historico.registrar(arquivo, _eventos())

        lidos = historico.ler(arquivo)

        assert len(lidos) == 4
        assert [e["resultado"] for e in lidos] == ["ok", "falha", "ok", "nao tentado"]

    def test_so_acrescenta(self, tmp_path):
        arquivo = tmp_path / "historico.jsonl"
        historico.registrar(arquivo, _eventos()[:1])
        historico.registrar(arquivo, _eventos()[1:])

        assert len(historico.ler(arquivo)) == 4

    def test_linha_corrompida_nao_derruba_a_leitura(self, tmp_path):
        arquivo = tmp_path / "historico.jsonl"
        historico.registrar(arquivo, _eventos()[:1])
        with arquivo.open("a", encoding="utf-8") as saida:
            saida.write("{isso nao e json\n")
        historico.registrar(arquivo, _eventos()[1:2])

        assert len(historico.ler(arquivo)) == 2

    def test_arquivo_inexistente_devolve_vazio(self, tmp_path):
        assert historico.ler(tmp_path / "nao-existe.jsonl") == []


class TestPlanilha:
    def test_uma_planilha_por_mes_com_as_tres_abas(self, tmp_path):
        arquivo = tmp_path / "historico.jsonl"
        historico.registrar(arquivo, _eventos())

        gravadas = historico.gerar_planilhas(arquivo, tmp_path / "logs")

        assert [p.name for p in gravadas] == [f"{historico._mes(_eventos()[0])}.xlsx"]
        wb = load_workbook(gravadas[0])
        assert wb.sheetnames == ["Envios", "Extracoes", "Rodadas"]

    def test_separa_envio_de_extracao(self, tmp_path):
        arquivo = tmp_path / "historico.jsonl"
        historico.registrar(arquivo, _eventos())

        gravadas = historico.gerar_planilhas(arquivo, tmp_path / "logs")

        wb = load_workbook(gravadas[0])
        assert wb["Extracoes"].max_row == 3  # cabecalho + 2
        assert wb["Envios"].max_row == 3

    def test_a_falha_aparece_com_o_motivo(self, tmp_path):
        """O ponto do log: o que NAO saiu precisa estar la, e dizer por que."""
        arquivo = tmp_path / "historico.jsonl"
        historico.registrar(arquivo, _eventos())

        gravadas = historico.gerar_planilhas(arquivo, tmp_path / "logs")

        wb = load_workbook(gravadas[0])
        extracoes = list(wb["Extracoes"].iter_rows(min_row=2, values_only=True))
        falha = [linha for linha in extracoes if linha[5] == "falha"]
        assert len(falha) == 1 and falha[0][6] == "timeout"

    def test_aba_rodadas_conta_por_dia_e_etapa(self, tmp_path):
        arquivo = tmp_path / "historico.jsonl"
        historico.registrar(arquivo, _eventos())

        gravadas = historico.gerar_planilhas(arquivo, tmp_path / "logs")

        wb = load_workbook(gravadas[0])
        por_etapa = {
            linha[1]: linha[2:]
            for linha in wb["Rodadas"].iter_rows(min_row=2, values_only=True)
        }
        assert por_etapa["extracao"] == (1, 1, 0, 2)      # ok, falha, nao tentado, total
        assert por_etapa["envio"] == (1, 0, 1, 2)

    def test_historico_vazio_nao_gera_planilha(self, tmp_path):
        assert historico.gerar_planilhas(tmp_path / "nada.jsonl", tmp_path / "logs") == []


class TestNaoConfundeComRelatorio:
    """A planilha mora dentro do FORMATADO_DIR, junto dos mapas."""

    def test_o_disparo_nao_a_le_como_relatorio(self, tmp_path):
        planilha = tmp_path / "Mapa de Estoque" / "logs" / "2026-09.xlsx"
        mapa = tmp_path / "Yuri Toso" / "MARJAN" / "2026-09" / "0112_MARJAN_2026-09.xlsx"

        assert _rotulo_do_arquivo(planilha) is None
        assert _rotulo_do_arquivo(mapa) == "2026-09"

    def test_o_nome_nao_entra_na_contagem_da_sequencia(self, tmp_path):
        """Por isso o nome nao pode comecar com digito+underscore."""
        assert _PADRAO_NUMERO.match("2026-09.xlsx") is None
        assert _PADRAO_NUMERO.match("0112_MARJAN_2026-09.xlsx") is not None

    def test_o_nome_gerado_obedece_a_regra(self, tmp_path):
        arquivo = tmp_path / "historico.jsonl"
        historico.registrar(arquivo, _eventos())

        gravadas = historico.gerar_planilhas(arquivo, tmp_path / "logs")

        assert _PADRAO_NUMERO.match(gravadas[0].name) is None


class TestNuncaDerrubaARodada:
    def test_destino_impossivel_nao_levanta(self, tmp_path):
        """Excel com o arquivo aberto e o caso real: a rodada apenas avisa e segue."""
        arquivo = tmp_path / "historico.jsonl"
        historico.registrar(arquivo, _eventos())
        ocupado = tmp_path / "ocupado"
        ocupado.write_text("sou um arquivo, nao uma pasta", encoding="utf-8")

        historico.atualizar(arquivo, ocupado, [])  # nao pode levantar

    def test_jsonl_em_caminho_invalido_nao_levanta(self, tmp_path):
        impossivel = tmp_path / "arquivo.txt" / "historico.jsonl"
        (tmp_path / "arquivo.txt").write_text("x", encoding="utf-8")

        historico.registrar(impossivel, _eventos())  # nao pode levantar

    def test_sem_eventos_nao_cria_arquivo(self, tmp_path):
        arquivo = tmp_path / "historico.jsonl"

        historico.registrar(arquivo, [])

        assert not arquivo.exists()
