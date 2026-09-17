"""Historico estruturado das rodadas, e a planilha que o torna legivel.

Por que um arquivo novo, e nao o logs/envios.jsonl que ja existe: aquele alimenta a COTA
horaria em src/disparo/limites.py, que conta uma linha por mensagem. Registrar falhas la
faria a cota contar mensagens que nunca sairam e apertar o limite sem motivo. Os dois
propositos sao diferentes — um e contador de defesa, este e memoria do que aconteceu.

O jsonl e a fonte da verdade, so acrescentando. A planilha e uma PROJECAO: se alguem deixar
o .xlsx aberto no Excel, o Windows trava a gravacao, e ai a rodada apenas avisa e segue —
nada se perde, porque o proximo comando reconstroi a planilha inteira a partir do jsonl.

Para refazer as planilhas a qualquer momento, sem rodar nada:

    python -m src.historico
"""

from __future__ import annotations

import json
import logging
from collections import Counter
from datetime import datetime
from pathlib import Path
from typing import Any, Iterable

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font
from openpyxl.utils import get_column_letter

log = logging.getLogger(__name__)

FONTE = "Arial"
TAMANHO = 9
LARGURA_MAXIMA = 60

COLUNAS_ENVIO = ["data", "hora", "comprador", "laboratorio", "periodo", "canal",
                 "destinatario", "resultado", "motivo", "id"]
COLUNAS_EXTRACAO = ["data", "hora", "comprador", "laboratorio", "periodo",
                    "resultado", "motivo", "arquivo"]
COLUNAS_RODADA = ["data", "etapa", "ok", "falha", "nao tentado", "total"]


def registrar(arquivo: Path, eventos: Iterable[dict[str, Any]]) -> None:
    """Acrescenta eventos ao jsonl. Falhar aqui nunca derruba a rodada.

    O registro e do que JA aconteceu: perder a anotacao e ruim, mas abortar depois de o
    mapa ter sido extraido e enviado seria pior.
    """
    eventos = list(eventos)
    if not eventos:
        return
    try:
        arquivo.parent.mkdir(parents=True, exist_ok=True)
        with arquivo.open("a", encoding="utf-8") as saida:
            for evento in eventos:
                saida.write(json.dumps(evento, ensure_ascii=False) + "\n")
    except OSError as erro:
        log.warning("Nao consegui gravar o historico em %s: %s", arquivo.name, erro)


def ler(arquivo: Path) -> list[dict[str, Any]]:
    """Le o jsonl inteiro. Linha corrompida e pulada, nao derruba a leitura."""
    if not arquivo.exists():
        return []
    eventos: list[dict[str, Any]] = []
    for linha in arquivo.read_text(encoding="utf-8").splitlines():
        if not linha.strip():
            continue
        try:
            registro = json.loads(linha)
        except ValueError:
            log.debug("Linha ilegivel em %s, ignorada", arquivo.name)
            continue
        if isinstance(registro, dict) and registro.get("em"):
            eventos.append(registro)
    return eventos


def evento_extracao(comprador: str, laboratorio: str, periodo: str, resultado: str,
                    motivo: str = "", arquivo: str = "") -> dict[str, Any]:
    """resultado: ok | falha | sem dados."""
    return {
        "em": datetime.now().isoformat(timespec="seconds"),
        "tipo": "extracao",
        "comprador": comprador,
        "laboratorio": laboratorio,
        "periodo": periodo,
        "resultado": resultado,
        "motivo": motivo,
        "arquivo": arquivo,
    }


def evento_envio(comprador: str, laboratorio: str, periodo: str, canal: str,
                 destinatario: str, resultado: str, motivo: str = "",
                 identificador: str = "") -> dict[str, Any]:
    """resultado: ok | falha | nao tentado."""
    return {
        "em": datetime.now().isoformat(timespec="seconds"),
        "tipo": "envio",
        "comprador": comprador,
        "laboratorio": laboratorio,
        "periodo": periodo,
        "canal": canal,
        "destinatario": destinatario,
        "resultado": resultado,
        "motivo": motivo,
        "id": identificador,
    }


def _mes(evento: dict[str, Any]) -> str:
    return str(evento.get("em", ""))[:7]


def _linha(evento: dict[str, Any], colunas: list[str]) -> list[Any]:
    em = str(evento.get("em", ""))
    valores: dict[str, Any] = {"data": em[:10], "hora": em[11:19]}
    for coluna in colunas:
        if coluna not in valores:
            valores[coluna] = evento.get(coluna, "")
    return [valores[coluna] for coluna in colunas]


def _aba(wb: Workbook, titulo: str, colunas: list[str], linhas: list[list[Any]]) -> None:
    ws = wb.create_sheet(titulo)
    ws.append([coluna.upper() for coluna in colunas])
    for linha in linhas:
        ws.append(linha)

    negrito = Font(name=FONTE, size=TAMANHO, bold=True)
    normal = Font(name=FONTE, size=TAMANHO)
    for celula in ws[1]:
        celula.font = negrito
        celula.alignment = Alignment(horizontal="center", vertical="center")
    for fila in ws.iter_rows(min_row=2):
        for celula in fila:
            celula.font = normal

    for indice, coluna in enumerate(colunas, start=1):
        conteudo = [len(str(linha[indice - 1] or "")) for linha in linhas]
        largura = max([len(coluna)] + conteudo, default=10)
        ws.column_dimensions[get_column_letter(indice)].width = min(largura + 2, LARGURA_MAXIMA)

    # Cabecalho congelado e autofiltro: e assim que o comprador acha a propria carteira
    # sem precisar de um arquivo separado por pessoa.
    ws.freeze_panes = "A2"
    if linhas:
        ws.auto_filter.ref = f"A1:{get_column_letter(len(colunas))}{len(linhas) + 1}"
    ws.sheet_view.showGridLines = False


def _resumo_por_dia(eventos: list[dict[str, Any]]) -> list[list[Any]]:
    """A aba Rodadas e derivada, nao registrada: assim nunca discorda das outras duas."""
    contagem: dict[tuple[str, str], Counter] = {}
    for evento in eventos:
        chave = (str(evento.get("em", ""))[:10], str(evento.get("tipo", "")))
        contagem.setdefault(chave, Counter())[str(evento.get("resultado", ""))] += 1
    linhas = []
    for (data, etapa), conta in sorted(contagem.items(), reverse=True):
        linhas.append([data, etapa, conta.get("ok", 0), conta.get("falha", 0),
                       conta.get("nao tentado", 0), sum(conta.values())])
    return linhas


def gerar_planilhas(historico: Path, destino: Path) -> list[Path]:
    """Reescreve uma planilha por mes em `destino`. Devolve as que conseguiu gravar.

    O nome NAO pode comecar com digito+underscore: src/sequencia.py conta qualquer arquivo
    assim como relatorio numerado, e o contador passaria a pular numeros. "2026-09.xlsx" e
    seguro, e o disparo tambem nao o confunde com relatorio (src/disparo/selecao.py).
    """
    eventos = ler(historico)
    if not eventos:
        return []

    gravadas: list[Path] = []
    for mes in sorted({_mes(e) for e in eventos if _mes(e)}):
        do_mes = [e for e in eventos if _mes(e) == mes]
        wb = Workbook()
        wb.remove(wb.active)
        _aba(wb, "Envios", COLUNAS_ENVIO,
             [_linha(e, COLUNAS_ENVIO) for e in do_mes if e.get("tipo") == "envio"])
        _aba(wb, "Extracoes", COLUNAS_EXTRACAO,
             [_linha(e, COLUNAS_EXTRACAO) for e in do_mes if e.get("tipo") == "extracao"])
        _aba(wb, "Rodadas", COLUNAS_RODADA, _resumo_por_dia(do_mes))

        arquivo = destino / f"{mes}.xlsx"
        try:
            destino.mkdir(parents=True, exist_ok=True)
            wb.save(arquivo)
            gravadas.append(arquivo)
        except OSError as erro:
            # Quase sempre o proprio Excel com o arquivo aberto. A planilha e projecao: o
            # jsonl ja tem o dado, e o proximo comando reconstroi tudo.
            log.warning(
                "Nao consegui gravar %s (%s). O historico esta salvo em %s e a planilha "
                "sera refeita na proxima rodada — feche o arquivo se estiver aberto.",
                arquivo.name, erro, historico.name,
            )
    return gravadas


def atualizar(historico: Path, destino: Path, eventos: Iterable[dict[str, Any]] = ()) -> None:
    """Registra os eventos e reescreve as planilhas. Nunca levanta excecao.

    Chamada no fim da extracao e do disparo, quando o trabalho ja terminou: nenhum problema
    de log pode mudar o codigo de saida de uma rodada que deu certo.
    """
    try:
        registrar(historico, eventos)
        gravadas = gerar_planilhas(historico, destino)
        if gravadas:
            log.info("Historico: %s", ", ".join(str(caminho) for caminho in gravadas))
    except Exception as erro:  # nenhum log vale derrubar uma rodada que ja terminou
        log.warning("Historico nao atualizado: %s", erro)
        log.debug("Detalhe", exc_info=True)


def main() -> int:
    from src.config import carregar_config
    from src.log import configurar as configurar_log

    configurar_log("execucao.log")
    cfg = carregar_config()
    gravadas = gerar_planilhas(cfg.historico_path, cfg.historico_dir)
    if not gravadas:
        log.info("Nada a gerar: %s esta vazio ou nao existe.", cfg.historico_path)
        return 0
    for caminho in gravadas:
        log.info("Gerado: %s", caminho)
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
